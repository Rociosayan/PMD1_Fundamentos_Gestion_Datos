# Genera 16_Lavanderia_PMD1.xlsx con el mismo formato de los 15 casos sinteticos
import numpy as np, pandas as pd, openpyxl
from openpyxl.styles import Font, PatternFill
rng=np.random.default_rng(20261022)
ciudades=["Lima","Arequipa","Trujillo","Piura","Cusco","Huancayo"]
cli=pd.DataFrame({"id_cliente":[f"C16{i:03d}" for i in range(1,111)],
 "nombre_cliente":[f"Cliente {i}" for i in range(1,111)],
 "ciudad":rng.choice(ciudades,110,p=[.45,.15,.12,.1,.1,.08]),
 "antiguedad_meses":rng.integers(1,37,110)})
for i in rng.choice(110,6,replace=False): cli.loc[i,"nombre_cliente"]="  "+cli.loc[i,"nombre_cliente"]+" "
tipos=[("lavado por kilo","basico",4.5),("lavado al seco","especial",9.0),("planchado","basico",3.5),("edredones","intermedio",7.0),("zapatillas","intermedio",6.0)]
cat=[]
for i in range(1,19):
    t,n,p=tipos[(i-1)%5]
    cat.append((f"P16{i:03d}",f"{t} {i}",n,round(p*rng.uniform(.85,1.2),2)))
cat=pd.DataFrame(cat,columns=["id_catalogo","nombre_item","nivel","precio_referencia_soles"])
sed=pd.DataFrame({"id_sede":[f"S16{i}" for i in range(1,6)],"nombre_sede":[f"Lavanderia sede {i}" for i in range(1,6)],"zona":["Centro","Norte","Sur","Este","Norte"]})
N=470; rows=[]
for k in range(N):
    c=rng.integers(0,18); t=cat.iloc[c]
    peso=round(float(np.clip(rng.gamma(3.2,2.2),1,25)),1)
    delic=int(min(rng.poisson(1.2+peso/8), 12))
    ocup=round(float(np.clip(rng.normal(55,20),5,100)),1)
    costo=t.precio_referencia_soles*peso*0.9+ 2.5*delic + 0.05*ocup + rng.normal(0,6)+ (8 if t.nivel=="especial" else 0)
    rows.append([f"O16{k+1:04d}",cli.id_cliente[rng.integers(0,110)],t.id_catalogo,sed.id_sede[rng.choice(5,p=[.3,.2,.2,.15,.15])],
      (pd.Timestamp("2025-01-01")+pd.Timedelta(days=int(rng.integers(0,365)))).strftime("%Y-%m-%d"),peso,delic,ocup,round(max(costo,6),2)])
op=pd.DataFrame(rows,columns=["id_operacion","id_cliente","id_catalogo","id_sede","fecha","peso_ropa_kg","prendas_delicadas","ocupacion_local_pct","costo_servicio_soles"])
op["peso_ropa_kg"]=op.peso_ropa_kg.astype(float); op["prendas_delicadas"]=op.prendas_delicadas.astype(float)
# errores didacticos
idx=rng.choice(N,22,replace=False)
op.loc[idx,"fecha"]=[pd.Timestamp(f).strftime("%d/%m/%Y").lstrip("0").replace("/0","/") for f in op.loc[idx,"fecha"]]
for col,n in [("peso_ropa_kg",4),("prendas_delicadas",5),("costo_servicio_soles",3),("ocupacion_local_pct",2)]:
    op.loc[rng.choice(N,n,replace=False),col]=np.nan
op.loc[rng.choice(N,2,replace=False),"peso_ropa_kg"]=[180.0,250.0]   # digitado en gramos/decenas: imposible
op.loc[rng.choice(N,1),"costo_servicio_soles"]=-35.0                 # costo negativo: error de registro
op.loc[rng.choice(N,2,replace=False),"id_catalogo"]="P16099"          # clave huerfana
dups=op.iloc[rng.choice(N,6,replace=False)]
op=pd.concat([op,dups]).sample(frac=1,random_state=7).reset_index(drop=True)
leeme=[("PMD1 — caso sintetico","16 Lavanderia"),("Origen","Datos inventados para practica; no son observaciones reales"),
("Unidad operativa","Una fila de Operaciones representa un pedido de lavado"),("Tablas","Operaciones, Clientes, Catalogo, Sedes"),
("Relaciones","Usar los identificadores compartidos; no hay tabla analitica preparada"),
("Calidad","Hay fechas con formatos distintos, nulos, filas repetidas, espacios en texto, claves sin correspondencia y valores fuera de rango"),
("Regresion","Se puede plantear una pregunta con una sola variable predictora"),("K-Means","Definir unidad de agrupacion, variables y escalado; interpretar segmentos"),
("Uso","Caso modelo creado por la docente para el ejemplo del PMD1"),("Semilla generadora","20261022; caso generado de forma reproducible"),
("Moneda","Soles peruanos (S/); todos los datos son sintéticos"),(None,None),
("Precios","precio_referencia_soles es una tarifa referencial por kg o por prenda; no son cotizaciones reales."),
("Unidades","peso_ropa_kg en kilogramos; ocupacion_local_pct es la ocupación de máquinas del local (0 a 100); costo_servicio_soles en soles."),
("Trabajo del alumno","Las tablas siguen separadas y conservan errores didácticos. El estudiante debe justificar limpieza, integración y modelos.")]
wb=openpyxl.Workbook(); ws=wb.active; ws.title="LEEME"
hdr=Font(bold=True,color="FFFFFFFF"); fill=PatternFill("solid",fgColor="FF164B70")
for r in leeme: ws.append(list(r))
for c in ws[1]: c.font=hdr; c.fill=fill
ws.column_dimensions["A"].width=27; ws.column_dimensions["B"].width=110
for name,df in [("Operaciones",op),("Clientes",cli),("Catalogo",cat),("Sedes",sed)]:
    w=wb.create_sheet(name); w.append(list(df.columns))
    for r in df.itertuples(index=False): w.append([None if (isinstance(v,float) and np.isnan(v)) else (v.item() if hasattr(v,"item") else v) for v in r])
    for c in w[1]: c.font=hdr; c.fill=fill
    for col in "ABCDEFGHI": w.column_dimensions[col].width=21
wb.save("16_Lavanderia_PMD1.xlsx"); print(op.shape, op.isna().sum().to_dict())
